"""
מחשבון יחסי מתכון — מחשב כמות ק"ג לכל רכיב לפי יחס המתכון.

שימוש:
  python recipe_ratio_calculator.py <קובץ_מתכונים.xlsx> <סה"כ_ק"ג> [פלט.xlsx]

עמודות נדרשות בקובץ הקלט:
  מתכון, רכיב, יחס

עמודה אופציונלית:
  סה"כ ק"ג   — אצווה ספציפית לכל מתכון (עוקפת את הארגומנט)
"""

import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

# ─── קבועים ────────────────────────────────────────────────────────────────
C_HEADER = "FF203864"
C_RECIPE = "FFD9E1F2"
C_ROW    = "FFFFFFFF"
C_ALT    = "FFF2F2F2"
C_TOTAL  = "FF92D050"

FONT_WHITE = Font(name="Arial", bold=True, color="FFFFFFFF", size=11)
FONT_BOLD  = Font(name="Arial", bold=True, size=10)
FONT_REG   = Font(name="Arial", size=10)

THIN   = Side(style="thin", color="FF000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center", wrap_text=True)

COL_RECIPE     = "מתכון"
COL_INGREDIENT = "רכיב"
COL_RATIO      = "יחס"
COL_TOTAL_KG   = 'סה"כ ק"ג'


# ─── פונקציות עזר ───────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _style_cell(cell, value, font=None, fill=None, alignment=None,
                border=True, number_format=None):
    cell.value = value
    if font:          cell.font = font
    if fill:          cell.fill = fill
    if alignment:     cell.alignment = alignment
    if border:        cell.border = BORDER
    if number_format: cell.number_format = number_format


def _merge_and_style(ws, row, col_start, col_end, value, font, fill,
                     alignment=None):
    ws.merge_cells(
        start_row=row, start_column=col_start,
        end_row=row, end_column=col_end,
    )
    cell = ws.cell(row=row, column=col_start)
    _style_cell(cell, value, font=font, fill=fill,
                alignment=alignment or CENTER)
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).border = BORDER


# ─── טעינה ──────────────────────────────────────────────────────────────────

def load_recipes(path: str) -> pd.DataFrame:
    df = pd.read_excel(path, dtype={COL_RECIPE: str, COL_INGREDIENT: str})
    required = [COL_RECIPE, COL_INGREDIENT, COL_RATIO]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"עמודות חסרות בקובץ: {missing}")
    df[COL_RATIO] = pd.to_numeric(df[COL_RATIO], errors="coerce").fillna(0)
    if COL_TOTAL_KG in df.columns:
        df[COL_TOTAL_KG] = pd.to_numeric(df[COL_TOTAL_KG],
                                          errors="coerce").fillna(0)
    return df


# ─── בניית Excel ─────────────────────────────────────────────────────────────

def build_excel(df: pd.DataFrame, total_kg_default: float,
                output_path: str) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "יחסי מתכון"
    ws.sheet_view.rightToLeft = True

    COLS = 5
    date_str = datetime.now().strftime("%d/%m/%Y")

    # שורה 1: כותרת ראשית
    _merge_and_style(ws, 1, 1, COLS,
                     f"מחשבון יחסי מתכון — {date_str}",
                     font=FONT_WHITE, fill=_fill(C_HEADER))
    ws.row_dimensions[1].height = 28

    # שורה 2: כותרות עמודות
    headers    = ["רכיב", "יחס", "אחוז", 'ק"ג', "הערה"]
    col_widths = [30, 12, 14, 14, 20]
    for i, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=2, column=i)
        _style_cell(cell, h, font=FONT_WHITE,
                    fill=_fill(C_HEADER), alignment=CENTER)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "A3"

    row = 3

    for recipe_name, grp in df.groupby(COL_RECIPE, sort=False):
        # קבע אצווה: מהעמודה בקובץ (שורה ראשונה) אחרת מהארגומנט
        total_kg = total_kg_default
        if COL_TOTAL_KG in grp.columns:
            val = grp[COL_TOTAL_KG].iloc[0]
            if val > 0:
                total_kg = val

        total_ratio = grp[COL_RATIO].sum()
        if total_ratio == 0:
            continue

        # כותרת מתכון
        _merge_and_style(
            ws, row, 1, COLS,
            f'◆  {recipe_name}   (אצווה: {total_kg:,.3f} ק"ג)',
            font=FONT_BOLD, fill=_fill(C_RECIPE), alignment=RIGHT,
        )
        ws.row_dimensions[row].height = 20
        row += 1

        for idx, (_, item) in enumerate(grp.iterrows()):
            ratio = item[COL_RATIO]
            pct   = ratio / total_ratio
            kg    = pct * total_kg
            fill  = _fill(C_ROW if idx % 2 == 0 else C_ALT)

            vals   = [item[COL_INGREDIENT], ratio, pct, kg, ""]
            fmts   = [None, "#,##0.###", "0.00%", '#,##0.000 ק"ג', None]
            fonts  = [FONT_REG, FONT_REG, FONT_REG, FONT_BOLD, FONT_REG]
            aligns = [RIGHT, CENTER, CENTER, CENTER, RIGHT]

            for col_i, (v, fmt, fnt, aln) in enumerate(
                zip(vals, fmts, fonts, aligns), start=1
            ):
                cell = ws.cell(row=row, column=col_i)
                _style_cell(cell, v, font=fnt, fill=fill,
                            alignment=aln, number_format=fmt)

            row += 1

        # שורת סה"כ
        _merge_and_style(ws, row, 1, 2, 'סה"כ',
                         font=FONT_BOLD, fill=_fill(C_TOTAL), alignment=RIGHT)
        cell_pct = ws.cell(row=row, column=3)
        _style_cell(cell_pct, 1.0, font=FONT_BOLD, fill=_fill(C_TOTAL),
                    alignment=CENTER, number_format="0.00%")
        cell_kg = ws.cell(row=row, column=4)
        _style_cell(cell_kg, total_kg, font=FONT_BOLD, fill=_fill(C_TOTAL),
                    alignment=CENTER, number_format='#,##0.000 ק"ג')
        ws.cell(row=row, column=5).border = BORDER
        ws.cell(row=row, column=5).fill = _fill(C_TOTAL)
        ws.row_dimensions[row].height = 18
        row += 2

    ws.auto_filter.ref = f"A2:{get_column_letter(COLS)}2"

    wb.save(output_path)
    return output_path


# ─── נקודת כניסה ─────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 3:
        print('שימוש: python recipe_ratio_calculator.py <מתכונים.xlsx> <ק"ג> [פלט.xlsx]')
        print()
        print("  עמודות נדרשות בקובץ: מתכון, רכיב, יחס")
        print('  עמודה אופציונלית:    סה"כ ק"ג  (אצווה ספציפית למתכון)')
        sys.exit(1)

    input_path  = sys.argv[1]
    total_kg    = float(sys.argv[2])
    output_path = sys.argv[3] if len(sys.argv) > 3 else (
        Path(input_path).stem
        + f"_יחסים_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )

    print(f"טוען: {input_path}")
    df = load_recipes(input_path)
    print(f"  מתכונים: {df[COL_RECIPE].nunique()}")
    print(f"  רכיבים:  {len(df)}")

    out = build_excel(df, total_kg, output_path)
    print(f'✅  נשמר: {out}')


if __name__ == "__main__":
    main()
