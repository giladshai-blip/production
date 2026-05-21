"""
תוכנית ייצור שבועית — מעבד קובץ WMS ומייצר גנט שבועי (ראשון–שישי).
"""

import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter
import pandas as pd

# ─── קבועים ────────────────────────────────────────────────────────────────
UNITS_PER_CART = 10

# פריסת עמודות: 4 קבועות + 6 ימים + סה"כ = 11
COL_FIXED     = 4
NUM_DAYS      = 6
COL_DAY_START = COL_FIXED + 1          # = 5
COL_TOTAL_COL = COL_FIXED + NUM_DAYS + 1  # = 11
TOTAL_COLS    = COL_TOTAL_COL

# צבעי רקע
C_ZERO      = "FFFF0000"   # אדום  — אין מלאי
C_LOW       = "FFFFA500"   # כתום  — מתחת למינימום
C_PLAN      = "FFFFFF00"   # צהוב  — תכנון
C_TOTAL     = "FF92D050"   # ירוק  — סה"כ משפחה
C_GRAND     = "FF4472C4"   # כחול  — סה"כ כולל
C_HEADER    = "FF203864"   # כחול כהה — כותרות
C_TODAY_HDR = "FF1F497D"   # כחול בהיר יותר — עמודת היום הנוכחי
C_CAT       = "FFD9E1F2"   # תכלת  — כותרת קטגוריה
C_EMPTY     = "FFFFFFFF"   # לבן   — תא ריק

FONT_WHITE = Font(name="Arial", bold=True, color="FFFFFFFF", size=11)
FONT_BOLD  = Font(name="Arial", bold=True, size=10)
FONT_REG   = Font(name="Arial", size=10)

THIN   = Side(style="thin", color="FF000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center", wrap_text=True)

# עמודות קובץ WMS
COL_FAM        = "משפחה"
COL_FAM_NAME   = "תאור משפחה"
COL_SKU        = "מק\"ט"
COL_DESC       = "תאור מוצר"
COL_STOCK      = "מלאי מרלוג ביח'"
COL_MIN        = "מינימום מלאי"
COL_DIFF_U     = "הפרש לייצור ביחידות"
COL_DIFF_B     = "הפרש לייצור באריזות"
COL_DAYS_STOCK = "ימי מלאי ממוצעים"

DAYS_HE = ["ראשון", "שני", "שלישי", "רביעי", "חמישי", "שישי"]


# ─── פונקציות עזר ───────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _style_cell(cell, value, font=None, fill=None, alignment=None, border=True, number_format=None):
    cell.value = value
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = BORDER
    if number_format:
        cell.number_format = number_format


def _merge_and_style(ws, row, col_start, col_end, value, font, fill, alignment=None):
    ws.merge_cells(
        start_row=row, start_column=col_start,
        end_row=row, end_column=col_end
    )
    cell = ws.cell(row=row, column=col_start)
    _style_cell(cell, value, font=font, fill=fill, alignment=alignment or CENTER)
    for c in range(col_start, col_end + 1):
        ws.cell(row=row, column=c).border = BORDER


def _day_index(days_stock: int) -> int:
    """ממיר ימי מלאי ממוצעים לאינדקס יום (0=ראשון … 5=שישי), עם הגבלה."""
    return min(max(int(days_stock), 0), 5)


# ─── טעינת נתונים ───────────────────────────────────────────────────────────

def load_wms(path: str) -> pd.DataFrame:
    df = pd.read_excel(path, dtype={COL_FAM: str, COL_SKU: str})
    required = [COL_FAM, COL_FAM_NAME, COL_SKU, COL_DESC, COL_DIFF_U, COL_DIFF_B]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"עמודות חסרות בקובץ WMS: {missing}")
    return df


def filter_and_prepare(df: pd.DataFrame) -> pd.DataFrame:
    df_need = df[df[COL_DIFF_U] < 0].copy()
    df_need["לייצור ביחידות"] = df_need[COL_DIFF_U].abs()
    df_need["לייצור באריזות"] = df_need[COL_DIFF_B].abs()
    if COL_STOCK not in df_need.columns:
        df_need[COL_STOCK] = 0
    if COL_MIN not in df_need.columns:
        df_need[COL_MIN] = 0
    if COL_DAYS_STOCK not in df_need.columns:
        df_need[COL_DAYS_STOCK] = 0
    df_need[COL_STOCK]      = pd.to_numeric(df_need[COL_STOCK],      errors="coerce").fillna(0)
    df_need[COL_MIN]        = pd.to_numeric(df_need[COL_MIN],        errors="coerce").fillna(0)
    df_need[COL_DAYS_STOCK] = pd.to_numeric(df_need[COL_DAYS_STOCK], errors="coerce").fillna(0).astype(int)
    return df_need


# ─── בניית Excel (גנט שבועי) ──────────────────────────────────────────────────

def build_excel(df_need: pd.DataFrame, output_path: str) -> str:
    today   = datetime.now()
    weekday = today.weekday()
    mapping = {6: 0, 0: 1, 1: 2, 2: 3, 3: 4, 4: 5, 5: 0}  # שבת → ראשון
    today_day_idx = mapping[weekday]
    date_str = today.strftime("%d/%m/%Y")
    week_num = today.isocalendar()[1]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"ייצור {today.strftime('%Y%m%d')}"
    ws.sheet_view.rightToLeft = True

    # ─── שורה 1: כותרת ראשית ───────────────────────────────────────────────
    title_text = f"תוכנית ייצור שבועית — שבוע {week_num}     {date_str}"
    _merge_and_style(ws, 1, 1, TOTAL_COLS, title_text,
                     font=FONT_WHITE, fill=_fill(C_HEADER))
    ws.row_dimensions[1].height = 28

    # ─── שורה 2: כותרות עמודות ─────────────────────────────────────────────
    col_widths = [22, 10, 32, 12, 10, 10, 10, 10, 10, 10, 12]
    fixed_headers = ["קטגוריה", "מק\"ט", "תאור מוצר", "מלאי נוכחי"]
    for i, (h, w) in enumerate(zip(fixed_headers, col_widths), start=1):
        _style_cell(ws.cell(row=2, column=i), h,
                    font=FONT_WHITE, fill=_fill(C_HEADER), alignment=CENTER)
        ws.column_dimensions[get_column_letter(i)].width = w

    for d in range(NUM_DAYS):
        col_i = COL_DAY_START + d
        hdr_fill = _fill(C_TODAY_HDR) if d == today_day_idx else _fill(C_HEADER)
        _style_cell(ws.cell(row=2, column=col_i), DAYS_HE[d],
                    font=FONT_WHITE, fill=hdr_fill, alignment=CENTER)
        ws.column_dimensions[get_column_letter(col_i)].width = col_widths[col_i - 1]

    _style_cell(ws.cell(row=2, column=COL_TOTAL_COL), "סה\"כ",
                font=FONT_WHITE, fill=_fill(C_HEADER), alignment=CENTER)
    ws.column_dimensions[get_column_letter(COL_TOTAL_COL)].width = col_widths[COL_TOTAL_COL - 1]
    ws.row_dimensions[2].height = 22

    row = 3
    grand_day_boxes = {d: 0 for d in range(NUM_DAYS)}
    grand_total = 0

    groups = df_need.groupby([COL_FAM, COL_FAM_NAME], sort=True)

    for (fam_code, fam_name), grp in groups:
        # ─── כותרת קטגוריה ─────────────────────────────────────────────────
        cat_text = f"◆  {fam_name}  ({fam_code})"
        _merge_and_style(ws, row, 1, TOTAL_COLS, cat_text,
                         font=Font(name="Arial", bold=True, size=10),
                         fill=_fill(C_CAT))
        ws.row_dimensions[row].height = 20
        row += 1

        fam_day_boxes = {d: 0 for d in range(NUM_DAYS)}
        fam_total = 0

        for _, item in grp.iterrows():
            stock   = int(item[COL_STOCK])
            min_s   = int(item[COL_MIN])
            boxes   = int(item["לייצור באריזות"])
            day_idx = _day_index(item[COL_DAYS_STOCK])

            if stock == 0:
                row_color = C_ZERO
            elif stock < min_s:
                row_color = C_LOW
            else:
                row_color = C_PLAN
            urgency_fill = _fill(row_color)

            # עמודות קבועות
            fixed_data = [
                (fam_name, None,    FONT_REG,  RIGHT),
                (item[COL_SKU],  None, FONT_REG,  CENTER),
                (item[COL_DESC], None, FONT_REG,  RIGHT),
                (stock,    "#,##0", FONT_BOLD, CENTER),
            ]
            for col_i, (v, fmt, fnt, aln) in enumerate(fixed_data, start=1):
                _style_cell(ws.cell(row=row, column=col_i), v,
                            font=fnt, fill=urgency_fill, alignment=aln, number_format=fmt)

            # שש עמודות ימים
            for d in range(NUM_DAYS):
                col_i = COL_DAY_START + d
                if d == day_idx:
                    _style_cell(ws.cell(row=row, column=col_i), boxes,
                                font=FONT_BOLD, fill=urgency_fill,
                                alignment=CENTER, number_format="#,##0")
                else:
                    _style_cell(ws.cell(row=row, column=col_i), None,
                                font=FONT_REG, fill=_fill(C_EMPTY), alignment=CENTER)

            # עמודת סה"כ שורה
            _style_cell(ws.cell(row=row, column=COL_TOTAL_COL), boxes,
                        font=FONT_BOLD, fill=urgency_fill,
                        alignment=CENTER, number_format="#,##0")

            fam_day_boxes[day_idx] += boxes
            fam_total += boxes
            grand_day_boxes[day_idx] += boxes
            grand_total += boxes
            row += 1

        # ─── שורת סה"כ משפחה ───────────────────────────────────────────────
        total_label = f"סה\"כ  {fam_name}"
        _merge_and_style(ws, row, 1, COL_FIXED, total_label,
                         font=FONT_BOLD, fill=_fill(C_TOTAL), alignment=RIGHT)
        for d in range(NUM_DAYS):
            col_i = COL_DAY_START + d
            val = fam_day_boxes[d] if fam_day_boxes[d] > 0 else ""
            _style_cell(ws.cell(row=row, column=col_i), val,
                        font=FONT_BOLD, fill=_fill(C_TOTAL),
                        alignment=CENTER, number_format="#,##0")
        _style_cell(ws.cell(row=row, column=COL_TOTAL_COL), fam_total,
                    font=FONT_BOLD, fill=_fill(C_TOTAL),
                    alignment=CENTER, number_format="#,##0")
        ws.row_dimensions[row].height = 18
        row += 1

    # ─── שורת סה"כ כולל ────────────────────────────────────────────────────
    _merge_and_style(ws, row, 1, COL_FIXED, "סה\"כ כולל",
                     font=FONT_WHITE, fill=_fill(C_GRAND), alignment=CENTER)
    for d in range(NUM_DAYS):
        col_i = COL_DAY_START + d
        val = grand_day_boxes[d] if grand_day_boxes[d] > 0 else ""
        _style_cell(ws.cell(row=row, column=col_i), val,
                    font=FONT_WHITE, fill=_fill(C_GRAND),
                    alignment=CENTER, number_format="#,##0")
    _style_cell(ws.cell(row=row, column=COL_TOTAL_COL), grand_total,
                font=FONT_WHITE, fill=_fill(C_GRAND),
                alignment=CENTER, number_format="#,##0")
    ws.row_dimensions[row].height = 22
    row += 1

    # ─── AutoFilter + הקפאת כותרות ─────────────────────────────────────────
    ws.auto_filter.ref = f"A2:{get_column_letter(TOTAL_COLS)}{row - 1}"
    ws.freeze_panes = "A3"

    wb.save(output_path)
    return output_path


# ─── נקודת כניסה ─────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print("שימוש: python daily_production.py <קובץ_WMS.xlsx> [קובץ_פלט.xlsx]")
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else (
        Path(input_path).stem + f"_גנט_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )

    print(f"טוען: {input_path}")
    df = load_wms(input_path)
    print(f"  סה\"כ שורות: {len(df)}")

    df_need = filter_and_prepare(df)
    print(f"  פריטים לייצור: {len(df_need)}")

    out = build_excel(df_need, output_path)
    print(f"✅  נשמר: {out}")


if __name__ == "__main__":
    main()
